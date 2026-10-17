import json, sys
from pathlib import Path
from types import SimpleNamespace
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pytest
import generate as g
from topics import parse_topics
from build_pptx import build

SRC = {"author": "Aliyev", "title": "Sud tibbiyoti", "year": 2020, "pages": "1-10"}


def _sources(tmp_path, with_text=True):
    t = tmp_path / "t.txt"
    t.write_text("matn " * 100 if with_text else "", encoding="utf-8")
    p = tmp_path / "s.json"
    p.write_text(json.dumps({"4-kurs-davolash-01": [{**SRC, "local_text_path": str(t)}]}), encoding="utf-8")
    return p


def _good():
    s = {"author": "Aliyev", "title": "Sud tibbiyoti", "page": "5"}
    return {"slides": [{"title": f"S{i}", "bullets": [{"text": "x", "source": s, "uncertain": i == 1}], "teacher_notes": "n"} for i in range(32)],
            "mcq": [{"question": "q", "options": list("abcd"), "correct_index": 1, "explanation": "e", "source": s} for _ in range(10)]}


def test_parse_topics():
    ts = parse_topics()
    assert len(ts) > 50 and ts[0]["id"] == "4-kurs-davolash-01"
    assert len({t["id"] for t in ts}) == len(ts)


def test_refuses_without_sources(tmp_path):
    class Boom:
        def __getattr__(self, n): raise AssertionError("API chaqirildi")
    r = g.run("4-kurs-davolash-01", tmp_path / "none.json", 0.5, False, tmp_path, client=Boom())
    assert r["status"] == "pending"
    r = g.run("4-kurs-davolash-01", _sources(tmp_path, False), 0.5, False, tmp_path, client=Boom())
    assert r["status"] == "pending"


def test_dry_run_and_budget(tmp_path, capsys):
    p = _sources(tmp_path)
    r = g.run("4-kurs-davolash-01", p, 0.5, True, tmp_path)
    assert r["status"] == "dry_run"
    assert g.est_cost(1_000_000, 1_000_000) == pytest.approx(g.PRICE_IN_PER_MTOK_USD + g.PRICE_OUT_PER_MTOK_USD)
    assert g.run("4-kurs-davolash-01", p, 0.0001, False, tmp_path, client=object())["status"] == "over_budget"


def test_mock_generate_and_pptx(tmp_path):
    calls = {}
    class C:
        class messages:
            @staticmethod
            def create(**kw):
                calls.update(kw)
                return SimpleNamespace(content=[SimpleNamespace(type="text", text=json.dumps(_good()))])
    r = g.run("4-kurs-davolash-01", _sources(tmp_path), 0.5, False, tmp_path, client=C())
    assert r["status"] == "ok" and calls["model"] == "claude-sonnet-5-5"
    assert calls["system"][0]["cache_control"] and calls["max_tokens"] == g.MAX_TOKENS
    data = json.load(open(r["path"], encoding="utf-8"))
    out = build(data, tmp_path / "o.pptx")
    from pptx import Presentation
    prs = Presentation(out)
    assert len(prs.slides) == 34 and prs.slides[1].has_notes_slide


def test_unsourced_bullet_rejected(tmp_path):
    d = _good(); d["slides"][0]["bullets"][0]["source"] = {"author": "Boshqa", "title": "X", "page": "1"}
    srcs = [{**SRC}]
    with pytest.raises(ValueError):
        g.validate_output(d, srcs)
