"""Inglizcha rejim (--lang en) mock testlari: haqiqiy API chaqirilmaydi."""
import json, sys
from pathlib import Path
from types import SimpleNamespace
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pytest
import generate as g
from build_pptx import build

SIM = {"author": "Payne-James J., Jones R., Karch S.B., Manlove J.", "title": "Simpson's Forensic Medicine, 13th ed.",
       "year": 2011, "pages": "Chapter 5, pp. 42-53"}
LOCAL = {"author": "G‘iyosov Z.A.", "title": "Sud tibbiyoti: darslik", "year": 2018, "pages": "Bob 8"}
TID = "4-kurs-davolash-01"


def _sources(tmp_path, local_first=True):
    a = tmp_path / "sim.txt"; a.write_text("[p. 42]\nPost-mortem interval text " * 50, encoding="utf-8")
    b = tmp_path / "giy.txt"; b.write_text("Murda sovishi " * 50, encoding="utf-8")
    es = [{**LOCAL, "local_text_path": str(b)}, {**SIM, "local_text_path": str(a)}]
    if not local_first:
        es.reverse()
    p = tmp_path / "s.json"; p.write_text(json.dumps({TID: es}), encoding="utf-8")
    return p


def _en_good():
    s = {"author": SIM["author"], "title": SIM["title"], "page": "45"}
    return {"slides": [{"title": f"Slide {i}", "bullets": [{"text": "Rigor mortis starts in the face and jaw.", "source": s, "uncertain": False}],
                        "teacher_notes": "10 min"} for i in range(32)],
            "mcq": [{"question": "q?", "options": list("abcd"), "correct_index": 2, "explanation": "e", "source": s} for _ in range(10)],
            "extras": [{"text": "Fact", "source": s}]}


class EnClient:
    prompts = []
    class messages:
        @staticmethod
        def create(**kw):
            txt = kw["messages"][0]["content"][1]["text"]
            EnClient.prompts.append((kw["system"][0]["text"], kw["messages"][0]["content"][0]["text"], txt))
            good = _en_good()
            if "STAGE 1" in txt:
                out = {"topic": "Forensic thanatology", "plan": [{"title": f"Slide {i}", "focus": "f", "diagram": False} for i in range(32)]}
            elif "STAGE 2" in txt:
                out = {"slides": good["slides"][:txt.count('"focus"')]}
            elif "STAGE 3.9" in txt:
                out = {"extras": good["extras"]}
            else:
                out = {"mcq": good["mcq"]}
            return SimpleNamespace(content=[SimpleNamespace(type="text", text=json.dumps(out))])


def test_en_sources_international_first(tmp_path):
    srcs = g.load_sources(TID, _sources(tmp_path), "en")
    assert g.is_international(srcs[0]) and not g.is_international(srcs[1])
    # uz rejimda tartib o'zgarmaydi
    assert not g.is_international(g.load_sources(TID, _sources(tmp_path))[0])
    only = g.load_sources(TID, _sources(tmp_path), "en", intl_only=True)
    assert len(only) == 1 and "Simpson" in only[0]["title"]


def test_en_prompt_is_english(tmp_path):
    srcs = g.load_sources(TID, _sources(tmp_path), "en")
    u = g.build_user_prompt({"title": "Sud-tibbiy tanatologiya", "course": "4 kurs"}, srcs, "en")
    assert "SOURCE 1 [INTERNATIONAL]" in u and "LOCAL (supplementary)" in u and "MANBA" not in u
    assert "ONLY in clear academic English" in g.SYSTEM_PROMPT_EN and "[p. 45]" in g.SYSTEM_PROMPT_EN
    assert "Faqat o'zbek tili" in g.SYSTEM_PROMPT  # uz default o'zgarmagan


def test_en_dry_run_and_pending(tmp_path):
    assert g.run(TID, _sources(tmp_path), 5.0, True, tmp_path, lang="en")["status"] == "dry_run"
    r = g.run(TID, tmp_path / "none.json", 5.0, False, tmp_path, client=object(), lang="en")
    assert r == {"status": "pending", "message": "LITERATURE PENDING"}
    with pytest.raises(ValueError):
        g.run(TID, _sources(tmp_path), 5.0, True, tmp_path, lang="fr")


def test_en_full_pipeline_and_pptx(tmp_path):
    EnClient.prompts.clear()
    r = g.run(TID, _sources(tmp_path), 5.0, False, tmp_path, client=EnClient(), lang="en")
    assert r["status"] == "ok" and r["path"].endswith(f"{TID}.en.json")
    assert not (tmp_path / f"{TID}.json").exists()  # uz fayli tegilmagan
    assert all(p[0] == g.SYSTEM_PROMPT_EN for p in EnClient.prompts)
    d = json.load(open(r["path"], encoding="utf-8"))
    assert d["lang"] == "en" and d["topic"] == "Forensic thanatology" and "Sud-tibbiy tanatologiya" in d["topic_uz"]
    assert d["sources"][0]["title"].startswith("Simpson")
    out = build(d, tmp_path / "o.en.pptx")
    from pptx import Presentation
    prs = Presentation(out)
    texts = [sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame]
    assert "References" in texts and "Quiz" in texts and "Question 1 / 40" in texts and "Answer 1: C" in texts
    assert any(t.startswith("Source: ") and "pp. 45" in t for t in texts)
    assert "Adabiyotlar" not in texts and not any("-bet" in t for t in texts)
    assert prs.slides[-1].shapes and "References" in [sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame]


def test_uz_pptx_unchanged(tmp_path):
    d = _en_good(); d.update(topic="Mavzu", sources=[SIM])
    prs_texts = lambda p: [sh.text_frame.text for sl in __import__("pptx").Presentation(p).slides for sh in sl.shapes if sh.has_text_frame]
    t = prs_texts(build(d, tmp_path / "u.pptx"))
    assert "Adabiyotlar" in t and "Bilimingizni sinang" in t and "Savol 1 / 10" in t and "Javob 1: C" in t


def test_en_rejects_cyrillic_and_bad_source():
    srcs = [{**SIM}]
    d = _en_good(); d["slides"][0]["bullets"][0]["text"] = "Трупное окоченение развивается раньше" * 3
    with pytest.raises(ValueError):
        g.validate_output(d, srcs, "en")
    d = _en_good(); d["slides"][0]["bullets"][0]["source"] = {"author": "Other", "title": "X", "page": "1"}
    with pytest.raises(ValueError):
        g.validate_output(d, srcs, "en")
    assert g.validate_output(_en_good(), srcs, "en")


def test_en_clean_text_keeps_apostrophe():
    d = g.clean_text({"slides": [{"title": "Simpson's guide", "bullets": [{"text": "dog's bite [Simpson, p. 5]"}]}]}, "en")
    assert d["slides"][0]["title"] == "Simpson's guide" and d["slides"][0]["bullets"][0]["text"] == "dog's bite"
    assert "‘" not in json.dumps(d, ensure_ascii=False)


def test_en_canonicalize_apostrophe_and_translit():
    """Run #22 (4-kurs-davolash-01:en) yiqilishi: model "G'iyosov Z.A." / qisqa sarlavha yozdi -> 'manbasiz band'."""
    import copy
    giy = {"author": "G‘iyosov Z.A.", "title": "Sud tibbiyoti (Tibbiyot oliy o‘quv yurtlari talabalari uchun darslik)", "year": 2018, "pages": "Bob 8"}
    rom = {"author": "Ромодановский П.О., Баринов Е.Х.", "title": "Судебная медицина в схемах и рисунках", "year": 2017, "pages": "2-modul"}
    pig = {"author": "Пиголкин Ю.И., Баринов Е.Х.", "title": "Судебная медицина: учебник", "year": 2002, "pages": "6-bob"}
    srcs = [giy, {**SIM}, rom, pig]
    d = _en_good()
    d["slides"][3]["bullets"][0]["source"] = {"author": "G'iyosov Z.A.", "title": "Sud tibbiyoti, Bob 8 (8.1)", "page": "217"}
    d["slides"][4]["bullets"][0]["source"] = {"author": "Romodanovsky P.O., Barinov E.Kh.", "title": "Forensic medicine in diagrams", "page": "41"}
    d["slides"][5]["bullets"][0]["source"] = {"author": "Pigolkin Yu.I. et al.", "title": "Forensic medicine textbook", "page": "70"}
    d["mcq"][0]["source"] = {"author": "Payne-James J, Jones R", "title": "Simpson's Forensic Medicine", "page": "42"}
    with pytest.raises(ValueError):
        g.validate_output(d, srcs, "en")
    uz = copy.deepcopy(d)
    g.canonicalize_sources(uz, srcs)  # uz rejim xulqi o'zgarmagan: fuzzy bog'lash yo'q
    assert uz["slides"][3]["bullets"][0]["source"]["author"] == "G'iyosov Z.A."
    g.canonicalize_sources(d, srcs, "en")
    assert d["slides"][3]["bullets"][0]["source"]["author"] == giy["author"]
    assert d["slides"][4]["bullets"][0]["source"]["title"] == rom["title"]
    assert d["slides"][5]["bullets"][0]["source"]["author"] == pig["author"]
    assert d["slides"][3]["bullets"][0]["source"]["page"] == "217"  # bet saqlanadi
    assert g.validate_output(d, srcs, "en")
    # noma'lum muallif baribir rad etiladi
    d["slides"][0]["bullets"][0]["source"] = {"author": "Knight B.", "title": "Forensic Pathology", "page": "1"}
    g.canonicalize_sources(d, srcs, "en")
    with pytest.raises(ValueError):
        g.validate_output(d, srcs, "en")


def test_en_prompt_requires_exact_source_fields():
    assert "EXACTLY as written in the SOURCE header" in g.SYSTEM_PROMPT_EN
    assert "EXACTLY" not in g.SYSTEM_PROMPT
