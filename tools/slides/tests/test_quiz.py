import json, re, sys
from pathlib import Path
from types import SimpleNamespace
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pytest
import generate_quiz as gq

SRC = {"author": "Aliyev", "title": "Sud tibbiyoti", "year": 2020, "pages": "1-10"}
TOPIC = "4-kurs-davolash-01"
WORDS = ["alfa", "beta", "gamma", "delta", "epsilon", "zeta", "eta", "teta", "yota", "kappa", "lambda", "myu", "nyu", "ksi", "omikron",
         "pi", "ro", "sigma", "tau", "upsilon", "fi", "xi", "psi", "omega", "bir", "ikki", "uch", "tort", "besh", "olti"]


def _sources(tmp_path):
    t = tmp_path / "t.txt"
    t.write_text("matn " * 100, encoding="utf-8")
    p = tmp_path / "s.json"
    p.write_text(json.dumps({TOPIC: [{**SRC, "local_text_path": str(t)}]}), encoding="utf-8")
    return p


def _uz_q(k, hard, ci=None):
    import random
    rnd = random.Random(k)
    stem = " ".join(rnd.choice(WORDS) + str(rnd.randint(0, 99)) for _ in range(9))
    return {"question": f"{stem}: COHb {20 + k}% da 1,5 soat ichida nima kuzatiladi?",
            "options": [f"{WORDS[(k + j) % 30]} variant {j} {k * 7 + j} mg" for j in range(4)],
            "correct_index": k % 4 if ci is None else ci,
            "explanation": f"Chunki {k + 3} foiz va COHb darajasi shuni ko‘rsatadi.",
            "difficulty": "hard" if hard else ("easy" if k % 2 else "medium"),
            "source": {"author": "Aliyev", "title": "Sud tibbiyoti", "page": str(k + 1)}}


class MockClient:
    """UZ BATCH va TARJIMA so'rovlariga javob beradi; qo'ng'iroqlarni yozib boradi."""
    def __init__(self, break_ru_once=False):
        self.calls, self.k, self.break_ru = [], 0, break_ru_once

    @property
    def messages(self):
        return self

    def create(self, **kw):
        self.calls.append(kw)
        txt = kw["messages"][0]["content"][-1]["text"]
        if txt.startswith("UZ BATCH"):
            m = re.search(r"Aynan (\d+) ta savol yozing: (\d+) ta oson.*? va (\d+) ta qiyin", txt)
            ne, nh = int(m.group(2)), int(m.group(3))
            qs = []
            for h in [False] * ne + [True] * nh:
                qs.append(_uz_q(self.k, h)); self.k += 1
            out = {"questions": qs}
        else:
            lang = re.search(r"TARJIMA (\w+)\.", txt).group(1)
            payload = json.loads(txt.split("INPUT_JSON:\n", 1)[1])
            res = []
            for q in payload["questions"]:
                r = {"id": q["id"], "question": f"[{lang}] " + q["question"], "options": [f"[{lang}] " + o for o in q["options"]],
                     "explanation": f"[{lang}] " + q["explanation"]}
                if self.break_ru and lang == "ru" and not getattr(self, "_broke", False):
                    r["question"] = r["question"].replace("20", "21"); self._broke = True
                res.append(r)
            out = {"questions": res}
            if "title" in payload:
                out["title"] = f"[{lang}] " + payload["title"]
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=json.dumps(out, ensure_ascii=False))], stop_reason="end_turn")


def _run(tmp_path, client, **kw):
    return gq.run(TOPIC, _sources(tmp_path), 5.0, False, tmp_path / "oq", client=client, **kw)


def test_numbers_and_abbrevs():
    assert gq.numbers("1,5 mg va 1.5; 20%") == {"1.5", "20"}
    assert gq.abbrevs("COHb GC-MS DNK Qaysi HbCO") == {"COHb", "GC-MS", "HbCO"}


def test_full_pipeline_structure(tmp_path):
    c = MockClient()
    r = _run(tmp_path, c)
    assert r["status"] == "ok"
    bank = json.load(open(r["path"], encoding="utf-8"))
    assert bank["topic_id"] == TOPIC and set(bank["title_i18n"]) == {"uz", "ru", "en"}
    qs = bank["questions"]
    assert len(qs) == 30
    assert sum(q["difficulty"] == "hard" for q in qs) == 15
    for q in qs:
        assert set(q["question_i18n"]) == set(q["options_i18n"]) == set(q["explanation_i18n"]) == {"uz", "ru", "en"}
        assert all(len(v) == 4 for v in q["options_i18n"].values())
        assert q["translation_status"] == {"uz": "source", "ru": "auto_unreviewed", "en": "auto_unreviewed"}
        assert set(q["source"]) == {"author", "title", "page"}
        # tarjima correct_index ni o'zgartirmaydi: yagona maydon
        assert 0 <= q["correct_index"] < 4
    assert bank["translation_status"]["ru"] == "auto_unreviewed"
    # 5 uz batch + 2 til x 3 chunk
    assert len(c.calls) == 5 + 6
    assert all(k["model"] == "claude-sonnet-5-5" for k in c.calls)
    assert all(k["system"][0]["cache_control"] for k in c.calls)
    assert c.calls[0]["messages"][0]["content"][0]["cache_control"]  # manba matni keshlanadi
    assert not (tmp_path / "oq" / f"{TOPIC}.unvalidated.json").exists()


def test_translation_retry_on_number_mismatch(tmp_path):
    c = MockClient(break_ru_once=True)
    r = _run(tmp_path, c)
    assert r["status"] == "ok"
    assert len(c.calls) == 5 + 6 + 1   # bitta qayta urinish
    assert "RAD ETILDI" in c.calls[-1]["messages"][0]["content"][-1]["text"] or any(
        "RAD ETILDI" in k["messages"][0]["content"][-1]["text"] for k in c.calls)


def test_translation_number_mismatch_fails_after_retry():
    uz = {"id": "x", "question": "20% COHb", "options": ["1", "2", "3", "4"], "explanation": "5 mg"}
    ok = {"question": "20% COHb", "options": ["1", "2", "3", "4"], "explanation": "5 mg"}
    gq.check_translation(uz, ok, "ru")
    for bad in ({**ok, "question": "21% COHb"}, {**ok, "options": ["1", "2", "3", "5"]}, {**ok, "explanation": "5,5 mg"},
                {**ok, "question": "20% COHB"}, {**ok, "options": ["1", "2", "3"]}):
        with pytest.raises(ValueError):
            gq.check_translation(uz, bad, "en")
    gq.check_translation({**uz, "explanation": "1,5 mg"}, {**ok, "explanation": "1.5 mg"}, "en")  # o'nli vergul = nuqta


def test_validate_uz_batch_rules():
    srcs = [dict(SRC)]
    good = [_uz_q(i, False) for i in range(6)]
    assert gq.validate_uz_batch(good, srcs, [], 6, 0, 6)
    with pytest.raises(ValueError):  # qiyinlik taqsimoti
        gq.validate_uz_batch(good, srcs, [], 3, 3, 6)
    bad = [dict(q) for q in good]; bad[0]["source"] = {"author": "Boshqa", "title": "X", "page": "1"}
    with pytest.raises(ValueError):
        gq.validate_uz_batch(bad, srcs, [], 6, 0, 6)
    bad = [dict(q) for q in good]; bad[1]["correct_index"] = 4
    with pytest.raises(ValueError):
        gq.validate_uz_batch(bad, srcs, [], 6, 0, 6)
    bad = [dict(q) for q in good]; bad[2]["options"] = bad[2]["options"][:3]
    with pytest.raises(ValueError):
        gq.validate_uz_batch(bad, srcs, [], 6, 0, 6)
    bad = [dict(q) for q in good]; bad[3]["source"] = {**bad[3]["source"], "page": ""}
    with pytest.raises(ValueError):
        gq.validate_uz_batch(bad, srcs, [], 6, 0, 6)
    with pytest.raises(ValueError):  # mavjud (slayd) MCQ bilan dublikat
        gq.validate_uz_batch(good, srcs, [gq._norm(good[0]["question"])], 6, 0, 6)


def test_validate_bank_30_rule(tmp_path):
    r = _run(tmp_path, MockClient())
    bank = json.load(open(r["path"], encoding="utf-8"))
    srcs = [dict(SRC)]
    assert gq.validate_bank(bank, srcs)
    short = json.loads(json.dumps(bank)); short["questions"].pop()
    with pytest.raises(ValueError):
        gq.validate_bank(short, srcs)
    diff = json.loads(json.dumps(bank)); diff["questions"][0]["difficulty"] = "hard"
    with pytest.raises(ValueError):
        gq.validate_bank(diff, srcs)
    nosrc = json.loads(json.dumps(bank)); nosrc["questions"][0]["source"]["author"] = "Boshqa"
    with pytest.raises(ValueError):
        gq.validate_bank(nosrc, srcs)
    badtr = json.loads(json.dumps(bank)); badtr["questions"][0]["options_i18n"]["en"][0] += " 99"
    with pytest.raises(ValueError):
        gq.validate_bank(badtr, srcs)
    badst = json.loads(json.dumps(bank)); badst["questions"][0]["translation_status"]["ru"] = "reviewed"
    with pytest.raises(ValueError):
        gq.validate_bank(badst, srcs)


def test_existing_mcq_used_for_dedupe(tmp_path):
    ex = tmp_path / "ex"; ex.mkdir()
    (ex / f"{TOPIC}.json").write_text(json.dumps({"mcq": [{"question": _uz_q(0, False)["question"]}]}), encoding="utf-8")
    assert len(gq.existing_mcq(TOPIC, ex)) == 1 and gq.existing_mcq("yoq", ex) == [] and gq.existing_mcq(TOPIC, None) == []
    c = MockClient()
    r = _run(tmp_path, c, existing_dir=ex)   # 1-urinish dublikat -> rad etiladi -> xato matni bilan qayta uriniladi
    assert r["status"] == "ok" and len(c.calls) == 5 + 6 + 1
    first, second = (k["messages"][0]["content"][-1]["text"] for k in c.calls[:2])
    assert _uz_q(0, False)["question"] in first          # mavjud savol promptga beriladi
    assert "RAD ETILDI" in second and "dublikat" in second
    bank = json.load(open(r["path"], encoding="utf-8"))
    assert all(q["question_i18n"]["uz"] != _uz_q(0, False)["question"] for q in bank["questions"])


def test_pending_dry_run_and_budget(tmp_path):
    class Boom:
        def __getattr__(self, n): raise AssertionError("API chaqirildi")
    assert gq.run(TOPIC, tmp_path / "none.json", 5.0, False, tmp_path, client=Boom())["status"] == "pending"
    p = _sources(tmp_path)
    assert gq.run(TOPIC, p, 5.0, True, tmp_path)["status"] == "dry_run"
    assert gq.run(TOPIC, p, 0.0001, False, tmp_path, client=Boom())["status"] == "over_budget"


def test_build_quiz_pptx_three_languages(tmp_path):
    from pptx import Presentation
    import build_quiz_pptx as bq
    r = _run(tmp_path, MockClient())
    bank = json.load(open(r["path"], encoding="utf-8"))
    paths = bq.build_all(bank, tmp_path)
    assert [Path(p).name for p in paths] == [f"{TOPIC}.uz.pptx", f"{TOPIC}.ru.pptx", f"{TOPIC}.en.pptx"]
    for p in paths:
        prs = Presentation(p)
        assert len(prs.slides) == 1 + 60 + 1
        assert prs.slide_width > prs.slide_height * 1.7  # 16:9
    ru = Presentation(paths[1])
    texts = " ".join(sh.text_frame.text for sl in ru.slides for sh in sl.shapes if sh.has_text_frame)
    assert "[ru]" in texts and "Ответ" in texts and "не проверен" in texts
