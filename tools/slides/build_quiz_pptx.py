"""Test banki JSON -> har til uchun alohida 16:9 .pptx (savol A-D, keyingi slaydda javob+tushuntirish+manba).
Uslub va yordamchilar build_pptx.py dan olinadi."""
import argparse, json, sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_pptx import (NAVY, GOLD, INK, GREY, WHITE, W, H, MX, _box, _text, _frame, _bullets)  # noqa: F401

LABELS = {
    "uz": {"sub": "Sinf testi", "q": "Savol", "a": "Javob", "src": "Manba", "page": "bet", "count": "ta test savoli",
           "refs": "Adabiyotlar", "tr": "", "easy": "oson-o‘rta", "hard": "qiyin"},
    "ru": {"sub": "Тест для аудитории", "q": "Вопрос", "a": "Ответ", "src": "Источник", "page": "стр.", "count": "вопросов",
           "refs": "Литература", "tr": "Автоматический перевод — не проверен", "easy": "лёгкий-средний", "hard": "сложный"},
    "en": {"sub": "Class quiz", "q": "Question", "a": "Answer", "src": "Source", "page": "p.", "count": "questions",
           "refs": "References", "tr": "Automatic translation — unreviewed", "easy": "easy-medium", "hard": "hard"},
}


def _src_line(q, L):
    s = q.get("source") or {}
    t = (s.get("title") or "").split("(")[0].split(",")[0].strip()
    return f"{L['src']}: {(s.get('author') or '').rstrip('.')}., {t}, {s.get('page', '')} {L['page']}"


def build(bank, lang, out_path):
    L = LABELS[lang]
    qs = bank["questions"]
    title = bank["title_i18n"][lang]
    foot = title[:80] + (f" · {L['tr']}" if L["tr"] else "")
    total = 1 + 2 * len(qs) + 1
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    n = 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s, 0, 0, W, H, NAVY); _box(s, MX, Inches(4.45), Inches(3), Inches(0.08), GOLD)
    _text(s, MX, Inches(1.5), W - 2 * MX, Inches(2.8), title, 38, WHITE, True, "Georgia", anchor=MSO_ANCHOR.BOTTOM)
    _text(s, MX, Inches(4.75), W - 2 * MX, Inches(0.6), f"{L['sub']} · {len(qs)} {L['count']}", 22, GOLD)
    if L["tr"]:
        _text(s, MX, Inches(5.5), W - 2 * MX, Inches(0.5), L["tr"], 16, GOLD)
    for i, q in enumerate(qs, 1):
        ci = q["correct_index"]
        opts = q["options_i18n"][lang][:4]
        qt = q["question_i18n"][lang]
        n += 1
        s = _frame(prs, f"{L['q']} {i} / {len(qs)}", n, total, foot)
        _text(s, MX, Inches(1.5), W - 2 * MX, Inches(1.7), qt, 24 if len(qt) < 170 else 20, INK, True)
        for k, opt in enumerate(opts):
            yy = Inches(3.3) + k * Inches(0.82)
            _box(s, MX, yy, Inches(0.7), Inches(0.68), GOLD, MSO_SHAPE.ROUNDED_RECTANGLE)
            _text(s, MX, yy, Inches(0.7), Inches(0.68), "ABCD"[k], 24, NAVY, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            _text(s, MX + Inches(0.9), yy, W - 2 * MX - Inches(0.9), Inches(0.68), opt, 20 if len(opt) < 80 else 16, INK, anchor=MSO_ANCHOR.MIDDLE)
        s.notes_slide.notes_text_frame.text = f"{q['id']} · {L['easy'] if q['difficulty'] != 'hard' else L['hard']}"
        n += 1
        s = _frame(prs, f"{L['a']} {i}: {'ABCD'[ci]}", n, total, foot)
        _text(s, MX, Inches(1.5), W - 2 * MX, Inches(1.1), f"{'ABCD'[ci]}. {opts[ci]}", 24 if len(opts[ci]) < 100 else 20, NAVY, True)
        ex = q["explanation_i18n"][lang]
        _text(s, MX, Inches(2.7), W - 2 * MX, Inches(3.6), ex, 20 if len(ex) < 420 else 17, INK)
        _text(s, MX, Inches(6.4), W - 2 * MX, Inches(0.5), _src_line(q, L), 11, GREY)
    n += 1
    s = _frame(prs, L["refs"], n, total, foot)
    _bullets(s, [{"text": f"{(x.get('author') or '').rstrip('.')}. {x.get('title')}. {x.get('year') or ''}. {x.get('pages')}"} for x in bank.get("sources", [])],
             MX, Inches(1.6), W - 2 * MX, Inches(5), 18)
    prs.save(out_path)
    return str(out_path)


def build_all(bank, out_dir, langs=("uz", "ru", "en")):
    out = []
    for lang in langs:
        out.append(build(bank, lang, Path(out_dir) / f"{bank['topic_id']}.{lang}.pptx"))
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("json_path")
    ap.add_argument("--langs", default="uz,ru,en")
    ap.add_argument("--out", default=None, help="chiqish papkasi (standart: JSON yonida)")
    a = ap.parse_args()
    bank = json.load(open(a.json_path, encoding="utf-8"))
    for p in build_all(bank, a.out or Path(a.json_path).parent, tuple(a.langs.split(","))):
        print(p)
