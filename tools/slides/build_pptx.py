"""JSON -> 16:9 .pptx (kontent vositasi). Katta ekran/AirPlay uchun yirik shrift, tahrirlanadigan shakllar."""
import json, math, re, sys
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

NAVY, GOLD, INK, GREY = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0xC9, 0xA2, 0x27), RGBColor(0x1F, 0x29, 0x37), RGBColor(0x6B, 0x72, 0x80)
PAPER, SOFT, WHITE = RGBColor(0xFB, 0xF8, 0xF1), RGBColor(0xE8, 0xEE, 0xF6), RGBColor(255, 255, 255)
W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.7)  # overscan uchun chekka


STR = {
    "uz": {"page": "-bet", "pages": "betlar", "source": "Manba", "subtitle": "Sud tibbiyoti · ma’ruza taqdimoti",
           "extras": "Qiziqarli ma’lumotlar", "quiz": "Bilimingizni sinang", "nq": "{n} ta test savoli",
           "q": "Savol {i} / {n}", "a": "Javob {i}: {c}", "refs": "Adabiyotlar", "check": "Tekshirish kerak (manbada noaniq)"},
    "en": {"page": "p.", "pages": "pp.", "source": "Source", "subtitle": "Forensic medicine · lecture presentation",
           "extras": "Did you know?", "quiz": "Quiz", "nq": "{n} quiz questions",
           "q": "Question {i} / {n}", "a": "Answer {i}: {c}", "refs": "References", "check": "Verify (uncertain in source)"},
}


def _box(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = fill; s.line.fill.background(); s.shadow.inherit = False
    return s


def _text(slide, x, y, w, h, text, size, color, bold=False, font="Calibri", align=None, anchor=None):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    if anchor: tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    if align: p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
    r.font.color.rgb = color
    return tf


def _src(b, lang="uz"):
    s = b.get("source") or {}
    t = (s.get("title") or "").split("(")[0].split(",")[0].strip()
    a = (s.get("author") or "").rstrip(".")
    if lang == "en":
        return f"{a}, {t}, p. {s.get('page','')}"
    return f"{a}., {t}, {s.get('page','')}-bet"


def _srcline(items, lang="uz"):
    by = {}
    for b in items:
        s = b.get("source") or {}
        k = ((s.get("author") or "").rstrip("."), (s.get("title") or "").split("(")[0].split(",")[0].strip())
        by.setdefault(k, [])
        p = str(s.get("page", "")).strip()
        if p and p not in by[k]:
            by[k].append(p)
    if lang == "en":
        return "; ".join(f"{a}, {t}, pp. {', '.join(ps)}" for (a, t), ps in by.items())
    return "; ".join(f"{a}., {t}, betlar: {', '.join(ps)}" for (a, t), ps in by.items())


def _frame(prs, title, n, total, topic):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s, 0, 0, W, H, PAPER)
    _box(s, 0, 0, W, Inches(1.25), NAVY)
    _box(s, 0, Inches(1.25), W, Inches(0.06), GOLD)
    _text(s, MX, Inches(0.12), W - 2 * MX, Inches(1.05), title, 30, WHITE, True, "Georgia", anchor=MSO_ANCHOR.MIDDLE)
    _text(s, MX, Inches(7.0), Inches(9), Inches(0.35), topic[:90], 11, GREY)
    _text(s, W - MX - Inches(1.5), Inches(7.0), Inches(1.5), Inches(0.35), f"{n} / {total}", 11, GREY, align=PP_ALIGN.RIGHT)
    return s


def _bullets(slide, items, x, y, w, h, size=22, gap=12):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    for i, b in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = "▸  " + b["text"]
        r.font.size, r.font.name = Pt(size), "Calibri"; r.font.color.rgb = INK
        p.space_after = Pt(gap)


def _diagram_basic(slide, d, x, y, w, h):
    """table / flow; noma'lum kind: header bo'lsa jadval, nodes bo'lsa flow, aks holda matn."""
    kind = d.get("kind")
    if d.get("header") and isinstance(d.get("header"), list):
        kind = "table"
    if kind == "table" and d.get("header"):
        rows = [d["header"]] + list(d.get("rows", []))[:6]
        cols = len(d["header"])
        fs = 17 if cols <= 2 else (15 if cols == 3 else 13)
        word = [max([len(wd) for r in rows if j < len(r) for wd in str(r[j]).split()] + [4]) for j in range(cols)]
        lens = [max(min(40, max([len(str(r[j])) if j < len(r) else 0 for r in rows] + [len(str(d["header"][j])) + 6, 16])), word[j] + 4) for j in range(cols)]
        need = [Inches(word[j] * fs * 0.62 / 72 + 0.3) for j in range(cols)]  # eng uzun so'z bo'linmasin
        tot = sum(lens)
        widths = [max(need[j], int(w * lens[j] / tot)) for j in range(cols)]
        over = sum(widths) - w
        if over > 0:  # ortiqchani keng ustunlardan olamiz
            slack = [widths[j] - need[j] for j in range(cols)]
            ss = sum(slack) or 1
            widths = [int(widths[j] - over * slack[j] / ss) for j in range(cols)]
        rh = min(Inches(0.95), int(h * 0.8 / len(rows)))
        gs = slide.shapes.add_table(len(rows), cols, x, y, w, rh * len(rows))
        t = gs.table
        for j in range(cols):
            t.columns[j].width = widths[j]
        for i, r in enumerate(rows):
            t.rows[i].height = rh
            for j in range(cols):
                c = t.cell(i, j); c.text = str(r[j]) if j < len(r) else ""
                c.margin_left = c.margin_right = Inches(0.08); c.margin_top = c.margin_bottom = Inches(0.04)
                c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (SOFT if i % 2 else WHITE)
                for p in c.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(fs + (1 if i == 0 else 0)); run.font.bold = i == 0
                        run.font.color.rgb = WHITE if i == 0 else INK
    elif d.get("nodes"):
        nodes = [str(n) for n in d.get("nodes", [])[:6]]
        gap = Inches(0.32); nh = min(Inches(0.85), (h - gap * (len(nodes) - 1)) / len(nodes))
        for i, n in enumerate(nodes):
            yy = y + i * (nh + gap)
            sh = _box(slide, x, yy, w, nh, NAVY if i % 2 == 0 else RGBColor(0x1F, 0x4E, 0x79), MSO_SHAPE.ROUNDED_RECTANGLE)
            tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.text = n
            for p in tf.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.size = Pt(16); run.font.color.rgb = WHITE; run.font.bold = True
            if i < len(nodes) - 1:
                a = _box(slide, x + w // 2 - Inches(0.15), yy + nh + Inches(0.02), Inches(0.3), gap - Inches(0.04), GOLD, MSO_SHAPE.DOWN_ARROW)
    else:
        _fallback_text(slide, d, x, y, w, h)
    if d.get("caption"):
        _text(slide, x, y + h + Inches(0.05), w, Inches(0.4), str(d["caption"])[:80], 11, GREY)


BLUE = RGBColor(0x1F, 0x4E, 0x79)
RICH = ("steps", "compare", "tree", "facts", "reaction", "cycle")
EMU_IN = 914400


def _est_lines(text, size, w_in):
    """Taxminiy qator soni (Calibri ~0.5 em/belgi)."""
    cpl = max(1, int(w_in * 72 / (size * 0.56)))
    return max(1, math.ceil(len(str(text)) / cpl))


def _fit(texts, w_in, h_in, hi, lo, spacing=1.25):
    """Matnlar (w_in x h_in) ga sig'adigan eng katta shrift (pt)."""
    for sz in range(int(hi), int(lo) - 1, -1):
        lines = sum(_est_lines(t, sz, w_in) for t in texts)
        if lines * sz * spacing / 72 <= h_in:
            return sz
    return lo


def _paras(sh, parts, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, margin=0.1, gap=2):
    """parts: har element — (matn,size,rang,bold) yoki shunday tuple'lar ro'yxati (bitta paragraf, bir necha run)."""
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin); tf.margin_top = tf.margin_bottom = Inches(0.04)
    for i, runs in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(gap)
        if isinstance(runs, tuple):
            runs = [runs]
        for t, sz, col, bold in runs:
            r = p.add_run(); r.text = str(t)
            r.font.size, r.font.bold, r.font.name = Pt(sz), bold, "Calibri"
            r.font.color.rgb = col


def _line(slide, x1, y1, x2, y2, color=GOLD, wpt=2.25, arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = color; c.line.width = Pt(wpt)
    if arrow:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
        te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def _card(slide, x, y, w, h, fill=WHITE, border=SOFT, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = border; s.line.width = Pt(1.25); s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def _strs(v, n):
    return [str(x) for x in v][:n] if isinstance(v, (list, tuple)) else []


def _d_steps(slide, d, x, y, w, h):
    items = [i for i in (d.get("items") or []) if isinstance(i, dict)][:6]
    if not items:
        return False
    cols = 1 if len(items) <= 3 else 2
    rows = math.ceil(len(items) / cols)
    gx, gy = Inches(0.18), Inches(0.14)
    cw, ch = (w - gx * (cols - 1)) / cols, (h - gy * (rows - 1)) / rows
    for k, it in enumerate(items):
        cx, cy = x + (k % cols) * (cw + gx), y + (k // cols) * (ch + gy)
        _card(slide, cx, cy, cw, ch)
        r = min(Inches(0.5), ch - Inches(0.2))
        c = _box(slide, cx + Inches(0.12), cy + (ch - r) / 2, r, r, NAVY, MSO_SHAPE.OVAL)
        _paras(c, [(str(k + 1), 16, WHITE, True)], PP_ALIGN.CENTER, margin=0)
        tx = cx + Inches(0.12) + r + Inches(0.08)
        tw = cw - (tx - cx) - Inches(0.08)
        title = re.sub(r"^\s*\d+\s*[.)]\s*", "", str(it.get("title", "")))
        text = str(it.get("text", ""))
        ts = _fit([title], tw / EMU_IN, ch / EMU_IN * 0.4, 16, 11)
        bs = _fit([text], tw / EMU_IN, ch / EMU_IN * 0.6, 14, 10)
        tb = slide.shapes.add_textbox(int(tx), int(cy), int(tw), int(ch))
        _paras(tb, [(title, ts, NAVY, True), (text, bs, INK, False)], margin=0.02)
    return True


def _d_compare(slide, d, x, y, w, h):
    cols = [c for c in (d.get("columns") or []) if isinstance(c, dict)][:3]
    if len(cols) < 2:
        return False
    gap = Inches(0.15); cw = (w - gap * (len(cols) - 1)) / len(cols)
    hh = Inches(0.6)
    for k, c in enumerate(cols):
        cx = x + k * (cw + gap)
        hd = _box(slide, cx, y, cw, hh, NAVY if k % 2 == 0 else BLUE, MSO_SHAPE.ROUNDED_RECTANGLE)
        ttl = str(c.get("title", ""))
        _paras(hd, [(ttl, _fit([ttl], cw / EMU_IN - 0.2, 0.5, 16, 11), WHITE, True)], PP_ALIGN.CENTER)
        body = _card(slide, cx, y + hh + Inches(0.06), cw, h - hh - Inches(0.06), SOFT, SOFT, MSO_SHAPE.RECTANGLE)
        items = _strs(c.get("items"), 5)
        wi, hi_ = cw / EMU_IN - 0.45, (h - hh) / EMU_IN - 0.2
        sz = _fit(items, wi, hi_ - 0.1 * len(items), 15, 10)
        _paras(body, [[("●  ", sz, GOLD, True), (t, sz, INK, False)] for t in items], anchor=MSO_ANCHOR.TOP, margin=0.1, gap=7)
    return True


def _d_tree(slide, d, x, y, w, h):
    kids = [k for k in (d.get("children") or []) if isinstance(k, dict)][:6]
    root = str(d.get("root") or "")
    if not kids or not root:
        return False
    rw = min(Inches(2.2), w * 0.34); gap = Inches(0.35)
    n = len(kids); vg = Inches(0.08)
    kh = (h - vg * (n - 1)) / n
    kx, kw = x + rw + gap, w - rw - gap
    rh = min(Inches(1.4), h)
    ry = y + (h - rh) / 2
    r = _box(slide, x, ry, rw, rh, NAVY, MSO_SHAPE.ROUNDED_RECTANGLE)
    _paras(r, [(root, _fit([root], rw / EMU_IN - 0.2, rh / EMU_IN - 0.2, 18, 11), WHITE, True)], PP_ALIGN.CENTER)
    bus = x + rw + gap / 2
    ys = [y + i * (kh + vg) + kh / 2 for i in range(n)]
    _line(slide, x + rw, ry + rh / 2, bus, ry + rh / 2, GOLD, 2)
    _line(slide, bus, ys[0], bus, ys[-1], GOLD, 2)
    for i, k in enumerate(kids):
        cy = y + i * (kh + vg)
        _line(slide, bus, ys[i], kx, ys[i], GOLD, 2)
        _card(slide, kx, cy, kw, kh, WHITE, GOLD)
        lab = str(k.get("label", "")); its = _strs(k.get("items"), 3)
        tb = slide.shapes.add_textbox(int(kx), int(cy), int(kw), int(kh))
        tw, th = kw / EMU_IN - 0.2, kh / EMU_IN - 0.08
        joined = "  ·  ".join(its)
        ls = 14 if its else 15
        isz = _fit([joined], tw, th - ls * 1.25 / 72, 12, 9) if its else 10
        _paras(tb, [(lab, ls, NAVY, True)] + ([(joined, isz, INK, False)] if its else []), margin=0.1, gap=0)
    return True


def _d_facts(slide, d, x, y, w, h):
    items = [i for i in (d.get("items") or []) if isinstance(i, dict)][:6]
    if not items:
        return False
    cols = 1 if len(items) == 1 else (3 if len(items) == 3 else 2)
    rows = math.ceil(len(items) / cols)
    gx, gy = Inches(0.18), Inches(0.16)
    cw, ch = (w - gx * (cols - 1)) / cols, min(Inches(1.7), (h - gy * (rows - 1)) / rows)
    for k, it in enumerate(items):
        cx, cy = x + (k % cols) * (cw + gx), y + (k // cols) * (ch + gy)
        _card(slide, cx, cy, cw, ch, WHITE, SOFT)
        _box(slide, cx, cy + Inches(0.1), Inches(0.07), ch - Inches(0.2), GOLD)
        lab, val, unit = str(it.get("label", "")), str(it.get("value", "")), str(it.get("unit") or "")
        tw = cw / EMU_IN - 0.3
        vs = _fit([val + (" " + unit if unit else "")], tw, ch / EMU_IN * 0.45, 30, 14, 1.15)
        ls = _fit([lab], tw, ch / EMU_IN * 0.4, 13, 10)
        runs = [(val, vs, NAVY, True)] + ([("  " + unit, max(11, int(vs * 0.5)), GREY, False)] if unit else [])
        tb = slide.shapes.add_textbox(int(cx + Inches(0.12)), int(cy), int(cw - Inches(0.12)), int(ch))
        _paras(tb, [(lab, ls, GREY, False), runs], margin=0.1, gap=1)
    return True


def _d_reaction(slide, d, x, y, w, h):
    rg, pr = _strs(d.get("reagents"), 3), _strs(d.get("products"), 3)
    if not rg or not pr:
        return False
    cond, obs = str(d.get("conditions") or ""), str(d.get("observation") or "")
    oh = Inches(0.8) if obs else 0
    ah = h - oh - (Inches(0.15) if obs else 0)
    aw = Inches(1.7); sw = (w - aw) / 2

    def side(items, sx, fill):
        n = len(items); pg = Inches(0.32 if n > 1 else 0)
        bh = min(Inches(1.0), (ah - pg * (n - 1)) / n)
        top = y + (ah - (bh * n + pg * (n - 1))) / 2
        for i, t in enumerate(items):
            by = top + i * (bh + pg)
            b = _box(slide, sx, by, sw, bh, fill, MSO_SHAPE.ROUNDED_RECTANGLE)
            _paras(b, [(t, _fit([t], sw / EMU_IN - 0.2, bh / EMU_IN - 0.1, 17, 10), WHITE, True)], PP_ALIGN.CENTER)
            if i < n - 1:
                _text(slide, sx, by + bh, sw, pg, "+", 16, GOLD, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    side(rg, x, NAVY)
    side(pr, x + sw + aw, BLUE)
    ay = y + ah / 2
    _line(slide, x + sw + Inches(0.1), ay, x + sw + aw - Inches(0.1), ay, GOLD, 4, arrow=True)
    if cond:
        _text(slide, x + sw, ay - Inches(0.85), aw, Inches(0.78), cond, _fit([cond], aw / EMU_IN - 0.1, 0.7, 13, 9), INK, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
    if obs:
        o = _card(slide, x, y + h - oh, w, oh, RGBColor(0xFF, 0xF6, 0xD6), GOLD)
        _paras(o, [[("Kuzatuv:  ", 13, NAVY, True), (obs, _fit([obs], w / EMU_IN - 1.8, oh / EMU_IN - 0.1, 15, 10), INK, False)]], margin=0.15)
    return True


def _d_cycle(slide, d, x, y, w, h):
    nodes = _strs(d.get("nodes"), 6)
    if len(nodes) < 3:
        return False
    n = len(nodes)
    nw, nh = min(Inches(2.2), w * 0.36), Inches(0.9)
    cx0, cy0 = x + w / 2, y + h / 2
    rx, ry = (w - nw) / 2, (h - nh) / 2
    ctr = []
    for i in range(n):
        a = -math.pi / 2 + 2 * math.pi * i / n
        ctr.append((cx0 + rx * math.cos(a), cy0 + ry * math.sin(a)))
    for i in range(n):
        (ax, ay), (bx, by) = ctr[i], ctr[(i + 1) % n]
        dx, dy = bx - ax, by - ay
        t = min((nw / 2 + Inches(0.06)) / abs(dx) if dx else 9, (nh / 2 + Inches(0.06)) / abs(dy) if dy else 9)
        if t < 0.5:
            _line(slide, ax + dx * t, ay + dy * t, bx - dx * t, by - dy * t, GOLD, 3, arrow=True)
    for i, (px, py) in enumerate(ctr):
        b = _box(slide, px - nw / 2, py - nh / 2, nw, nh, NAVY if i % 2 == 0 else BLUE, MSO_SHAPE.ROUNDED_RECTANGLE)
        _paras(b, [(nodes[i], _fit([nodes[i]], nw / EMU_IN - 0.2, nh / EMU_IN - 0.1, 14, 10), WHITE, True)], PP_ALIGN.CENTER)
    return True


_RENDER = {"steps": _d_steps, "compare": _d_compare, "tree": _d_tree, "facts": _d_facts, "reaction": _d_reaction, "cycle": _d_cycle}


def _flat_texts(v, out, depth=0):
    if depth > 4:
        return
    if isinstance(v, str):
        if v.strip():
            out.append(v.strip())
    elif isinstance(v, dict):
        for k, x in v.items():
            if k not in ("kind", "caption"):
                _flat_texts(x, out, depth + 1)
    elif isinstance(v, (list, tuple)):
        for x in v:
            _flat_texts(x, out, depth + 1)


def _diagram_has_content(d):
    if not isinstance(d, dict):
        return False
    t = []
    _flat_texts(d, t)
    return bool(t)


def _fallback_text(slide, d, x, y, w, h):
    """Noma'lum/yaroqsiz diagramma: matn ro'yxati sifatida (crash yo'q)."""
    t = []
    _flat_texts(d, t)
    t = t[:10]
    if not t:
        return
    sz = _fit(t, w / EMU_IN - 0.5, h / EMU_IN - 0.2 - 0.08 * len(t), 16, 10)
    b = _card(slide, x, y, w, h, SOFT, SOFT, MSO_SHAPE.RECTANGLE)
    _paras(b, [[("▸  ", sz, GOLD, True), (s, sz, INK, False)] for s in t], anchor=MSO_ANCHOR.TOP, gap=5, margin=0.15)


def _diagram(slide, d, x, y, w, h):
    fn = _RENDER.get(d.get("kind"))
    if not fn:
        return _diagram_basic(slide, d, x, y, w, h)
    cap = str(d.get("caption") or "")[:80]
    ch = Inches(0.4) if cap else 0
    try:
        ok = fn(slide, d, x, y, w, h - ch)
    except Exception:  # yaroqsiz maydon turi — matn fallback
        ok = False
    if not ok:
        _fallback_text(slide, d, x, y, w, h - ch)
    if cap:
        _text(slide, x, y + h - ch + Inches(0.05), w, ch, cap, 12, GREY)


def _clean_ref(t):
    """Ichki izohlar (egasi taqdim etgan, yopiq foydalanish, PDF-bet) slaydga chiqmasin."""
    t = re.sub(r"\s*\((?:egasi taqdim etgan|yopiq foydalanish)[^)]*\)", "", t)
    t = re.sub(r";?\s*(?:egasi taqdim etgan|yopiq foydalanish)[^.;)]*", "", t)
    t = re.sub(r"\s*\(?\b\d+-qism,\s*PDF-bet[^)]*\)?", "", t)
    t = re.sub(r"\s*\(PDF-bet[^)]*\)", "", t)
    return re.sub(r"\.\s*\.", ".", t).strip()


def build(data, out_path, lang=None):
    lang = lang or data.get("lang") or "uz"
    L = STR.get(lang, STR["uz"])
    prs = Presentation(); prs.slide_width, prs.slide_height = W, H
    topic = data.get("topic", "")
    # Test savollari (mcq) prezentatsiyaga qo'yilmaydi: ular faqat ilovadagi
    # alohida «Sinfda test» rejimida (egasi qarori, 2026-10-18).
    slides, extras = data["slides"], data.get("extras", [])
    total = 1 + len(slides) + (1 if extras else 0) + 1
    n = 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s, 0, 0, W, H, NAVY); _box(s, MX, Inches(4.45), Inches(3), Inches(0.08), GOLD)
    _text(s, MX, Inches(1.9), W - 2 * MX, Inches(2.4), topic, 40, WHITE, True, "Georgia", anchor=MSO_ANCHOR.BOTTOM)
    _text(s, MX, Inches(4.75), W - 2 * MX, Inches(0.6), L["subtitle"], 20, GOLD)
    for sl in slides:
        n += 1
        s = _frame(prs, sl["title"], n, total, topic)
        d = sl.get("diagram")
        has_d = _diagram_has_content(d)
        rich = has_d and d.get("kind") in RICH
        bw = Inches(5.0) if rich else (Inches(5.9) if has_d else W - 2 * MX)
        tot_len = sum(len(x["text"]) for x in sl["bullets"])
        size = (22 if tot_len < 600 else 20) if not has_d else (18 if tot_len < 450 else 16)
        _bullets(s, sl["bullets"], MX, Inches(1.55), bw, Inches(4.9), size, 8 if has_d else 12)
        if has_d:
            if rich:
                _diagram(s, d, MX + Inches(5.25), Inches(1.55), W - 2 * MX - Inches(5.25), Inches(4.8))
            else:
                _diagram(s, d, MX + Inches(6.15), Inches(1.6), W - 2 * MX - Inches(6.15), Inches(4.2))
        _text(s, MX, Inches(6.45), W - 2 * MX, Inches(0.5), L["source"] + ": " + _srcline(sl["bullets"], lang), 13, GREY)
        unsure = [b["text"] for b in sl["bullets"] if b.get("uncertain")]
        notes = sl.get("teacher_notes", "")
        if unsure:
            notes += "\n\n" + L.get("check", "Check") + ":\n- " + "\n- ".join(unsure)
        s.notes_slide.notes_text_frame.text = notes
    if extras:
        n += 1
        s = _frame(prs, L["extras"], n, total, topic)
        _bullets(s, [{"text": e["text"]} for e in extras], MX, Inches(1.6), W - 2 * MX, Inches(5.2), 18)
    n += 1
    s = _frame(prs, L["refs"], n, total, topic)
    refs = {}
    for x in data.get("sources", []):
        k = ((x.get("author") or "").rstrip("."), x.get("title") or "", x.get("year") or "")
        refs.setdefault(k, [])
        if x.get("pages") and x["pages"] not in refs[k]:
            refs[k].append(str(x["pages"]))
    items = [{"text": _clean_ref(f"{a}. {t}. {y}. {'; '.join(ps)}")} for (a, t, y), ps in refs.items()]
    _bullets(s, items, MX, Inches(1.6), W - 2 * MX, Inches(5.2), 20 if len(items) <= 4 else (16 if len(items) <= 7 else 13), 8)
    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    d = json.load(open(sys.argv[1], encoding="utf-8"))
    print(build(d, sys.argv[2] if len(sys.argv) > 2 else sys.argv[1].replace(".json", ".pptx")))
