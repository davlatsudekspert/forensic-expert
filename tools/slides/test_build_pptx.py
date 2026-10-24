"""build_pptx: v2 diagram turlari (DECK_SPEC_V2 §2) — limitdagi eng uzun matnlar bilan.
Ishga tushirish: python3 -m pytest tools/slides/test_build_pptx.py
LibreOffice + pdftoppm bo'lsa PDF/PNG render ham tekshiriladi (SLIDES_PNG_DIR bilan PNG saqlanadi)."""
import os, shutil, subprocess, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import pytest
from pptx import Presentation
from build_pptx import build

SRC = {"author": "Yuldashev Z.A.", "title": "Toksikologik kimyo", "page": "22"}


def L(n, c="Ш"):
    return (c * n)[:n]


def LONG(n):
    return ("Uzun matn " * 20)[:n]


DIAGRAMS = {
    "table": {"kind": "table", "header": [LONG(12)] * 4, "rows": [[LONG(30)] * 4 for _ in range(6)], "caption": LONG(80)},
    "flow": {"kind": "flow", "nodes": [LONG(40)] * 6, "caption": LONG(80)},
    "steps": {"kind": "steps", "items": [{"title": LONG(30), "text": LONG(80)}] * 6, "caption": LONG(80)},
    "compare2": {"kind": "compare", "columns": [{"title": LONG(24), "items": [LONG(60)] * 5}] * 2, "caption": LONG(80)},
    "compare3": {"kind": "compare", "columns": [{"title": LONG(24), "items": [LONG(60)] * 5}] * 3},
    "tree": {"kind": "tree", "root": LONG(30), "children": [{"label": LONG(26), "items": [LONG(40)] * 3}] * 6, "caption": LONG(80)},
    "facts": {"kind": "facts", "items": [{"label": LONG(24), "value": "W" * 20, "unit": "mg/kg" + "x" * 7}] * 6, "caption": LONG(80)},
    "reaction": {"kind": "reaction", "reagents": [LONG(40)] * 3, "conditions": LONG(40), "products": [LONG(40)] * 3,
                 "observation": LONG(60), "caption": LONG(80)},
    "cycle": {"kind": "cycle", "nodes": [LONG(28)] * 6, "caption": LONG(80)},
    "cycle3": {"kind": "cycle", "nodes": ["A", "B", "C"]},
}


def _slide(title, d):
    return {"title": title[:70], "bullets": [{"text": LONG(160), "source": SRC} for _ in range(5)],
            "diagram": d, "teacher_notes": "n"}


def _deck(slides):
    return {"topic": "Test mavzu", "slides": slides, "extras": [], "sources": [{"author": "Y", "title": "T", "year": 2020, "pages": "1"}]}


def _count(p):
    return len(Presentation(str(p)).slides)


@pytest.mark.parametrize("name", list(DIAGRAMS))
def test_each_kind_max_limits(tmp_path, name):
    out = build(_deck([_slide(name, DIAGRAMS[name])]), tmp_path / "a.pptx")
    assert _count(out) == 1 + 1 + 1  # titul + 1 slayd + adabiyotlar


def test_table_shows_six_rows(tmp_path):
    out = build(_deck([_slide("t", DIAGRAMS["table"])]), tmp_path / "a.pptx")
    tbl = [sh for sh in Presentation(str(out)).slides[1].shapes if sh.has_table][0].table
    assert len(tbl.rows) == 7  # sarlavha + 6


@pytest.mark.parametrize("d", [
    {"kind": "hologram", "items": ["a", "b"], "foo": {"bar": "baz"}},
    {"kind": "steps", "items": "notalist"},
    {"kind": "tree", "root": "r", "children": [None, 5]},
    {"kind": "cycle", "nodes": ["only"]},
    {"kind": "reaction", "reagents": [], "products": []},
    {"kind": "compare", "columns": [{"title": "x"}]},
    {"kind": "facts", "items": [{"label": 1, "value": 2}, "bad"]},
    {"kind": "table", "header": ["a"], "rows": [["x", "y"]]},
    {"kind": "flow", "nodes": [1, 2]},
    {}, None, "string", {"kind": "steps"},
])
def test_unknown_or_bad_no_crash(tmp_path, d):
    out = build(_deck([_slide("x", d)]), tmp_path / "a.pptx")
    assert _count(out) == 3


def test_all_kinds_deck_count(tmp_path):
    sl = [_slide(k, d) for k, d in DIAGRAMS.items()]
    out = build(_deck(sl), tmp_path / "a.pptx")
    assert _count(out) == len(sl) + 2


@pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftoppm")), reason="LibreOffice/pdftoppm yo'q")
def test_render_pdf_png(tmp_path):
    sl = [_slide(k, d) for k, d in DIAGRAMS.items()]
    out = build(_deck(sl), tmp_path / "a.pptx")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), str(out)],
                   check=True, capture_output=True, timeout=180)
    pdf = tmp_path / "a.pdf"
    assert pdf.exists()
    dest = Path(os.environ.get("SLIDES_PNG_DIR", tmp_path))
    dest.mkdir(parents=True, exist_ok=True)
    subprocess.run(["pdftoppm", "-r", "60", "-png", str(pdf), str(dest / "limits")], check=True, timeout=120)
    assert len(list(dest.glob("limits-*.png"))) == len(sl) + 2
